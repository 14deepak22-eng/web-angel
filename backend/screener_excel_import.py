"""
Imports fundamentals from Screener.in's per-company Excel export — this is
the file you download from a company's page on screener.in via "Export to
Excel". It's a completely different format from a screen's CSV export: one
file per company, with 10 years of P&L/Balance Sheet/Cash Flow data on the
"Data Sheet" tab (the other tabs are just formula-driven views of the same
data, so we read "Data Sheet" directly).

WHAT THIS DERIVES vs WHAT IT CAN'T:
  Derivable directly from the numbers in this file: revenue & EPS CAGR (3Y/5Y),
  OPM and margin change, ROE, ROCE, Debt/Equity, Interest Coverage, CFO, PAT,
  latest QoQ revenue growth, P/E, P/B, EV/EBITDA, Dividend Yield, PEG.

  NOT available in this file (Screener's free export doesn't include these) —
  left blank, need manual entry or the CSV/manual route instead:
    - Promoter holding % and pledge % (shareholding pattern isn't in this export)
    - Current ratio (no current-vs-non-current asset/liability split given)
    - Sector P/E, historical median P/E, fair value estimate (need external data)
    - True FCF (no explicit capex line — we approximate FCF as
      CFO + Cash from Investing Activity, which is a rough stand-in, not a
      textbook FCF calculation, and is flagged as such in the warnings)
"""

import io
import openpyxl


def _read_section(rows: list, start_idx: int) -> tuple[dict, tuple | None]:
    """
    Reads labeled rows following a section header (e.g. 'PROFIT & LOSS') until
    a blank-label row is hit. Returns ({label: (values across years...)}, dates).
    """
    data = {}
    dates = None
    i = start_idx
    while i < len(rows):
        row = rows[i]
        label = row[0]
        if label is None and i != start_idx:
            break
        if label == "Report Date":
            dates = row[1:]
        elif label is not None and i != start_idx:
            data[label] = row[1:]
        i += 1
    return data, dates


def _cagr(values: tuple, years: int) -> float | None:
    if len(values) <= years:
        return None
    start, end = values[-1 - years], values[-1]
    if start is None or end is None or start <= 0:
        return None
    return round(((end / start) ** (1 / years) - 1) * 100, 2)


def parse_screener_excel(file_bytes: bytes) -> dict:
    """
    Returns {"error": str|None, "company_name": str, "fundamentals": {...},
    "valuation": {...}, "warnings": [...]}
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        return {"error": f"Could not read this as an Excel file: {e}"}

    if "Data Sheet" not in wb.sheetnames:
        return {"error": "This doesn't look like a Screener.in company export "
                          "(expected a 'Data Sheet' tab). Make sure you used "
                          "'Export to Excel' from a company's page, not a screen's CSV export."}

    ws = wb["Data Sheet"]
    rows = list(ws.iter_rows(values_only=True))

    section_start = {}
    for i, row in enumerate(rows):
        if row[0] in ("PROFIT & LOSS", "Quarters", "BALANCE SHEET", "CASH FLOW:", "PRICE:", "DERIVED:"):
            section_start[row[0]] = i

    required = ["PROFIT & LOSS", "BALANCE SHEET", "CASH FLOW:", "DERIVED:"]
    missing = [s for s in required if s not in section_start]
    if missing:
        return {"error": f"Expected sections not found in this file: {', '.join(missing)}. "
                          f"The export format may have changed — check the file manually."}

    try:
        company_name = rows[0][1]
        current_price = rows[7][1]
        market_cap = rows[8][1]

        pl, _ = _read_section(rows, section_start["PROFIT & LOSS"])
        bs, _ = _read_section(rows, section_start["BALANCE SHEET"])
        cf, _ = _read_section(rows, section_start["CASH FLOW:"])
        derived, _ = _read_section(rows, section_start["DERIVED:"])
        q, _ = _read_section(rows, section_start["Quarters"]) if "Quarters" in section_start else ({}, None)

        warnings = []

        opex_labels = ["Raw Material Cost", "Change in Inventory", "Power and Fuel",
                       "Other Mfr. Exp", "Employee Cost", "Selling and admin", "Other Expenses"]
        n_years = len(pl["Sales"])
        op_profit = []
        for i in range(n_years):
            expenses = sum((pl.get(l, [0] * n_years)[i] or 0) for l in opex_labels)
            op_profit.append(pl["Sales"][i] - expenses)
        opm = [(op / sales * 100) if sales else None for op, sales in zip(op_profit, pl["Sales"])]

        rev_cagr3 = _cagr(pl["Sales"], 3)
        rev_cagr5 = _cagr(pl["Sales"], 5)

        adj_shares = derived.get("Adjusted Equity Shares in Cr")
        eps_cagr3 = eps_cagr5 = eps_latest = None
        if adj_shares:
            eps_series = [np / sh for np, sh in zip(pl["Net profit"], adj_shares) if sh]
            eps_cagr3 = _cagr(tuple(eps_series), 3)
            eps_cagr5 = _cagr(tuple(eps_series), 5)
            eps_latest = eps_series[-1] if eps_series else None
        else:
            warnings.append("Couldn't find share count data — EPS-based metrics (EPS CAGR, P/E, PEG) skipped.")

        equity = [e + r for e, r in zip(bs["Equity Share Capital"], bs["Reserves"])]
        pat_latest = pl["Net profit"][-1]
        roe = round(pat_latest / equity[-1] * 100, 2) if equity[-1] else None

        ebit_latest = pl["Profit before tax"][-1] + pl["Interest"][-1]
        capital_employed = equity[-1] + bs["Borrowings"][-1]
        roce = round(ebit_latest / capital_employed * 100, 2) if capital_employed else None

        de = round(bs["Borrowings"][-1] / equity[-1], 3) if equity[-1] else None
        int_cov = round(ebit_latest / pl["Interest"][-1], 2) if pl["Interest"][-1] else None

        cfo_latest = cf["Cash from Operating Activity"][-1]
        fcf_approx = None
        if "Cash from Investing Activity" in cf:
            fcf_approx = round(cfo_latest + cf["Cash from Investing Activity"][-1], 1)
            warnings.append("fcf is an approximation (CFO + Cash from Investing Activity), "
                             "not a true CFO-minus-capex calculation — this export doesn't "
                             "give a separate capex line.")

        qoq_rev = None
        if q and "Sales" in q and len(q["Sales"]) >= 2 and q["Sales"][-2]:
            qoq_rev = round((q["Sales"][-1] / q["Sales"][-2] - 1) * 100, 2)

        margin_chg_bps = None
        if len(opm) >= 2 and opm[-1] is not None and opm[-2] is not None:
            margin_chg_bps = round((opm[-1] - opm[-2]) * 100, 1)

        net_margin = round(pat_latest / pl["Sales"][-1] * 100, 2) if pl["Sales"][-1] else None

        pe = pb = ev_ebitda = div_yield = peg = None
        if eps_latest and current_price:
            pe = round(current_price / eps_latest, 2)
            if eps_cagr3 and eps_cagr3 > 0:
                peg = round(pe / eps_cagr3, 2)
        if adj_shares and current_price and equity[-1]:
            bvps = equity[-1] / adj_shares[-1]
            pb = round(current_price / bvps, 2) if bvps else None
        if market_cap and op_profit[-1]:
            ev = market_cap + bs["Borrowings"][-1] - bs.get("Cash & Bank", [0] * n_years)[-1]
            ev_ebitda = round(ev / op_profit[-1], 2)
        if "Dividend Amount" in pl and market_cap:
            div_yield = round(pl["Dividend Amount"][-1] / market_cap * 100, 3)

        fundamentals = {
            "roe": roe, "roce": roce, "ebitdaMargin": round(opm[-1], 2) if opm[-1] is not None else None,
            "marginChgBps": margin_chg_bps, "netMargin": net_margin,
            "revCagr3": rev_cagr3, "revCagr5": rev_cagr5,
            "epsCagr3": eps_cagr3, "epsCagr5": eps_cagr5, "qoqRev": qoq_rev,
            "de": de, "intCov": int_cov, "cfo": cfo_latest, "pat": pat_latest, "fcf": fcf_approx,
        }
        fundamentals = {k: v for k, v in fundamentals.items() if v is not None}

        valuation = {
            "cmp": current_price, "pe": pe, "pb": pb, "evEbitda": ev_ebitda,
            "divYield": div_yield, "peg": peg,
        }
        valuation = {k: v for k, v in valuation.items() if v is not None}

        warnings.append("Not available from this export — still need manual entry: "
                         "promoterHold, promoterPledge, currentRatio, sectorPE, medianPE, fairValue, fcfYield")

        return {
            "error": None,
            "company_name": company_name,
            "fundamentals": fundamentals,
            "valuation": valuation,
            "warnings": warnings,
        }

    except (KeyError, IndexError, TypeError) as e:
        return {"error": f"This file's layout didn't match what was expected (missing or "
                          f"differently-placed data: {e}). Screener occasionally changes "
                          f"their export format — this parser may need updating."}
